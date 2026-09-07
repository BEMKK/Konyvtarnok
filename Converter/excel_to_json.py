import json
import openpyxl
import wx


def convert_excel_to_json():
    # 1. wxApp inicializálása (szükséges a párbeszédablakokhoz)
    app = wx.App(False)

    file_path = "Enekeskonyvek_adatai.xlsx"

    try:
        # Excel fájl betöltése
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        # Fejlécek beolvasása
        headers = [cell.value for cell in sheet[1] if cell.value is not None]
        total_rows = sheet.max_row - 1  # Fejléc nélküli sorok száma

        if total_rows <= 0:
            wx.MessageBox(
                "Az Excel fájl nem tartalmaz adatokat!",
                "Hiba",
                wx.OK | wx.ICON_ERROR,
            )
            return

        # 2. Folyamatjelző ablak (ProgressDialog) létrehozása
        progress_dlg = wx.ProgressDialog(
            "Konvertálás folyamatban",
            "Excel adatok feldolgozása...",
            maximum=total_rows,
            style=wx.PD_APP_MODAL | wx.PD_AUTO_HIDE | wx.PD_CAN_ABORT,
        )

        data = []
        for index, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=1
        ):
            # Mégse gomb ellenőrzése
            keep_going, _ = progress_dlg.Update(
                index, f"Sorok feldolgozása: {index} / {total_rows}"
            )
            if not keep_going:
                progress_dlg.Destroy()
                wx.MessageBox(
                    "A folyamatot megszakítottad.",
                    "Megszakítva",
                    wx.OK | wx.ICON_WARNING,
                )
                return

            if any(row):
                row_dict = {headers[i]: row[i] for i in range(len(headers))}
                data.append(row_dict)

        progress_dlg.Destroy()

        # Mentés JSON-ba
        output_filename = "enekeskonyvek_adatai.json"
        with open(output_filename, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        # 3. Sikeres lefutás megerősítő ablak
        wx.MessageBox(
            f"A JSON adatbázis sikeresen frissítve!\n\nFeldolgozott sorok száma: {len(data)}",
            "Sikeres",
            wx.OK | wx.ICON_INFORMATION,
        )

    except FileNotFoundError:
        wx.MessageBox(
            f"A(z) '{file_path}' fájl nem található!",
            "Hiányzó fájl",
            wx.OK | wx.ICON_ERROR,
        )
    except Exception as e:
        wx.MessageBox(
            f"Hiba történt a konverzió során:\n\n{e}",
            "Hiba",
            wx.OK | wx.ICON_ERROR,
        )


if __name__ == "__main__":
    convert_excel_to_json()